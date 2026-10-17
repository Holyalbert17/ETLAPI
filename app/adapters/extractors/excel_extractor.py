import re
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.domain.schemas import ColumnInfo, SourceSchema, TableInfo
from app.ports.base_extractor import BaseExtractor
from app.services.identifiers import validate_ident

_UNSAFE = re.compile(r"[^A-Za-z0-9_]+")


def safe_name(raw: object, *, fallback: str) -> str:
    text = "" if raw is None else str(raw).strip()
    cleaned = _UNSAFE.sub("_", text).strip("_")
    if not cleaned or cleaned[0].isdigit():
        cleaned = f"{fallback}_{cleaned}" if cleaned else fallback
    return validate_ident(cleaned, what="columna de Excel")


def _python_type(value: object) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, (float, Decimal)):
        return "number"
    return "string"


class ExcelExtractor(BaseExtractor):
    """
    Lee .xlsx. Cada hoja es una tabla.

    config: path, sheets (opcional), header_row (1-based, default 1).
    """

    def connect(self) -> None:
        path = self.config.get("path")
        if not path:
            raise ValueError("Excel requiere 'path'")
        self.workbook = load_workbook(path, data_only=True)
        wanted = self.config.get("sheets")
        allowed = {str(name) for name in wanted} if wanted else None
        self.sheets: dict[str, str] = {}
        for sheet in self.workbook.worksheets:
            if allowed is not None and sheet.title not in allowed:
                continue
            table = safe_name(sheet.title, fallback="sheet")
            if table in self.sheets:
                raise ValueError(f"Dos hojas resuelven al mismo nombre de tabla: {table}")
            self.sheets[table] = sheet.title
        if not self.sheets:
            raise ValueError("El libro no tiene hojas utilizables")

    def close(self) -> None:
        workbook = getattr(self, "workbook", None)
        if workbook is not None:
            workbook.close()

    def _matrix(self, sheet_title: str) -> tuple[list[str], list[tuple]]:
        sheet = self.workbook[sheet_title]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        header_index = max(int(self.config.get("header_row", 1)) - 1, 0)
        if len(rows) <= header_index:
            raise ValueError(f"La hoja {sheet_title} no tiene fila de encabezado")
        names: list[str] = []
        seen: dict[str, int] = {}
        for index, cell in enumerate(rows[header_index]):
            name = safe_name(cell if cell not in (None, "") else None, fallback=f"col_{index + 1}")
            if name in seen:
                seen[name] += 1
                name = safe_name(f"{name}_{seen[name]}", fallback=f"col_{index + 1}")
            else:
                seen[name] = 1
            names.append(name)
        body = []
        for row in rows[header_index + 1 :]:
            values = tuple(row[: len(names)])
            if any(value is not None and value != "" for value in values):
                body.append(values)
        return names, body

    def discover_schema(self) -> SourceSchema:
        if not hasattr(self, "workbook"):
            self.connect()
        tables: list[TableInfo] = []
        for table, sheet_title in self.sheets.items():
            headers, body = self._matrix(sheet_title)
            kinds = {name: "string" for name in headers}
            for row in body[:50]:
                for name, value in zip(headers, row, strict=False):
                    if value is not None and value != "" and kinds[name] == "string":
                        kinds[name] = _python_type(value)
            columns = [
                ColumnInfo(name=name, type=kinds[name], nullable=True, primary_key=(name == "id"))
                for name in headers
            ]
            tables.append(TableInfo(name=table, columns=columns))
        return SourceSchema(tables=tables)

    def extract(
        self,
        table: str,
        batch_size: int = 500,
        checkpoint: dict | None = None,
    ) -> Iterator[dict]:
        del batch_size
        if not hasattr(self, "workbook"):
            self.connect()
        sheet_title = self.sheets.get(table)
        if sheet_title is None:
            raise ValueError(f"Hoja no encontrada: {table}")
        headers, body = self._matrix(sheet_title)
        skip = int((checkpoint or {}).get("skip") or 0)
        for row in body[skip:]:
            record = {}
            for name, value in zip(headers, row, strict=False):
                record[name] = None if value == "" else value
            self.bytes_read += len(repr(record))
            yield record
