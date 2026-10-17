import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]


def _workbook(path, sheets: dict[str, list[list]]):
    book = Workbook()
    first = True
    for title, rows in sheets.items():
        sheet = book.active if first else book.create_sheet(title)
        sheet.title = title
        first = False
        for row in rows:
            sheet.append(row)
    book.save(path)
    book.close()


def test_excel_two_sheets_and_second_source(client, admin_headers, tmp_path):
    ventas = tmp_path / "ventas.xlsx"
    _workbook(
        ventas,
        {
            "clientes": [
                ["id", "nombre", "correo", "alta"],
                [1, "Ana", "ana@example.com", datetime(2024, 1, 15)],
                [2, "Beto", "no-es-email", datetime(2024, 2, 2)],
                [None, None, None, None],
                [3, "Cia", "cia@example.com", datetime(2024, 3, 3)],
            ],
            "pedidos": [
                ["id", "sku", "total"],
                [10, "A-1", 1200.5],
                [11, "B-2", 99],
            ],
        },
    )
    created = client.post(
        "/api/v1/sources",
        headers=admin_headers,
        json={
            "name": "ventas_xlsx",
            "type": "excel",
            "config": {
                "path": str(ventas),
                "constraints": {"correo": {"format": "email"}},
            },
            "dedup_key": ["id"],
            "dedup_policy": "upsert",
        },
    )
    assert created.status_code == 201, created.text
    source_id = created.json()["id"]

    schema = client.get(f"/api/v1/sources/{source_id}/schema", headers=admin_headers)
    assert schema.status_code == 200, schema.text
    tables = {table["name"]: [col["name"] for col in table["columns"]] for table in schema.json()["tables"]}
    assert tables["clientes"] == ["id", "nombre", "correo", "alta"]
    assert tables["pedidos"] == ["id", "sku", "total"]
    assert schema.json()["diff"]["added_tables"] == ["clientes", "pedidos"]

    started = client.post(f"/api/v1/extract/{source_id}", headers=admin_headers)
    job = client.get(f"/api/v1/jobs/{started.json()['id']}", headers=admin_headers).json()
    assert job["status"] == "partial"
    assert job["rows_inserted"] == 4
    assert job["errors_count"] >= 1

    clients = client.get("/api/v1/data/ventas_xlsx/clientes?order_by=id", headers=admin_headers)
    assert clients.status_code == 200, clients.text
    rows = clients.json()["data"]
    assert [row["id"] for row in rows] == [1, 3]
    assert rows[0]["alta"] == "2024-01-15T00:00:00"
    assert rows[0]["nombre"] == "Ana"

    orders = client.get("/api/v1/data/ventas_xlsx/pedidos", headers=admin_headers).json()
    assert orders["total"] == 2
    assert {row["sku"] for row in orders["data"]} == {"A-1", "B-2"}

    again = client.post(f"/api/v1/extract/{source_id}", headers=admin_headers)
    second = client.get(f"/api/v1/jobs/{again.json()['id']}", headers=admin_headers).json()
    assert second["rows_updated"] == 4
    assert client.get("/api/v1/data/ventas_xlsx/clientes", headers=admin_headers).json()["total"] == 2

    catalogo = tmp_path / "catalogo.xlsx"
    _workbook(
        catalogo,
        {
            "Hoja 1": [
                ["sku", "stock", "precio lista"],
                ["ZX-9", 4, "$ 15.00"],
                ["ZX-8", 0, "€ 8,50"],
            ]
        },
    )
    other = client.post(
        "/api/v1/sources",
        headers=admin_headers,
        json={
            "name": "catalogo_xlsx",
            "type": "excel",
            "config": {"path": str(catalogo)},
            "dedup_key": ["sku"],
        },
    )
    assert other.status_code == 201, other.text
    client.post(f"/api/v1/extract/{other.json()['id']}", headers=admin_headers)
    data = client.get("/api/v1/data/catalogo_xlsx/Hoja_1", headers=admin_headers)
    assert data.status_code == 200, data.text
    assert data.json()["total"] == 2
    prices = {row["sku"]: row["precio_lista"] for row in data.json()["data"]}
    assert prices["ZX-9"] == "15.00"
    assert prices["ZX-8"] == "8.50"


def test_json_file_end_to_end(client, admin_headers, tmp_path):
    path = tmp_path / "ciudades.json"
    path.write_text(
        json.dumps({"items": [{"id": "1", "ciudad": "CABA"}, {"id": "2", "ciudad": "Rosario"}]}),
        encoding="utf-8",
    )
    created = client.post(
        "/api/v1/sources",
        headers=admin_headers,
        json={
            "name": "ciudades_json",
            "type": "json",
            "config": {"path": str(path), "records_path": "items"},
            "dedup_key": ["id"],
        },
    )
    assert created.status_code == 201, created.text
    source_id = created.json()["id"]
    schema = client.get(f"/api/v1/sources/{source_id}/schema", headers=admin_headers)
    assert schema.status_code == 200, schema.text
    names = [column["name"] for column in schema.json()["tables"][0]["columns"]]
    assert names == ["id", "ciudad"]

    client.post(f"/api/v1/extract/{source_id}", headers=admin_headers)
    data = client.get("/api/v1/data/ciudades_json/default?filter=ciudad:Rosario", headers=admin_headers)
    assert data.status_code == 200, data.text
    assert data.json()["total"] == 1
    assert data.json()["data"][0]["ciudad"] == "Rosario"


def test_example_excel_and_json_files(client, admin_headers):
    created = client.post(
        "/api/v1/sources",
        headers=admin_headers,
        json={
            "name": "ventas_ejemplo",
            "type": "excel",
            "config": {"path": str(ROOT / "examples" / "ventas.xlsx"), "sheets": ["clientes", "pedidos"]},
            "canonical_mapping": {"clientes": {"nombre": "full_name", "correo": "email"}},
            "dedup_key": ["id"],
            "dedup_policy": "upsert",
            "contains_pii": True,
        },
    )
    assert created.status_code == 201, created.text
    source_id = created.json()["id"]
    started = client.post(f"/api/v1/extract/{source_id}", headers=admin_headers)
    job = client.get(f"/api/v1/jobs/{started.json()['id']}", headers=admin_headers).json()
    assert job["status"] == "success"
    assert job["rows_inserted"] == 4

    people = client.get("/api/v1/data/ventas_ejemplo/clientes?order_by=id", headers=admin_headers).json()
    assert people["total"] == 2
    assert {row["full_name"] for row in people["data"]} == {"Ana López", "Beto Ruiz"}
    assert {row["email"] for row in people["data"]} == {"ana@example.com", "beto@example.com"}

    orders = client.get("/api/v1/data/ventas_ejemplo/pedidos", headers=admin_headers).json()
    assert orders["total"] == 2
    assert {row["sku"] for row in orders["data"]} == {"A-1", "B-2"}

    cities = client.post(
        "/api/v1/sources",
        headers=admin_headers,
        json={
            "name": "ciudades_ejemplo",
            "type": "json",
            "config": {
                "path": str(ROOT / "examples" / "ciudades.json"),
                "records_path": "items",
                "table": "ciudades",
            },
            "dedup_key": ["id"],
        },
    )
    assert cities.status_code == 201, cities.text
    client.post(f"/api/v1/extract/{cities.json()['id']}", headers=admin_headers)
    data = client.get(
        "/api/v1/data/ciudades_ejemplo/ciudades?filter=ciudad:Córdoba",
        headers=admin_headers,
    )
    assert data.status_code == 200, data.text
    assert data.json()["total"] == 1
    assert data.json()["data"][0]["ciudad"] == "Córdoba"


def test_excel_rejects_bad_books_and_sanitizes_headers(tmp_path):
    from datetime import date
    from decimal import Decimal

    from app.adapters.extractors.excel_extractor import ExcelExtractor, _python_type, safe_name

    assert safe_name("1 total", fallback="col_1") == "col_1_1_total"
    assert safe_name(None, fallback="col_2") == "col_2"
    assert _python_type(True) == "boolean"
    assert _python_type(date(2024, 1, 1)) == "date"
    assert _python_type(Decimal("1.5")) == "number"

    missing = ExcelExtractor({})
    try:
        missing.connect()
        raise AssertionError("debía exigir path")
    except ValueError as exc:
        assert "path" in str(exc)

    blank = tmp_path / "vacio.xlsx"
    _workbook(blank, {"datos": []})
    empty = ExcelExtractor({"path": str(blank), "header_row": 3})
    empty.connect()
    try:
        empty.discover_schema()
        raise AssertionError("debía fallar sin encabezado")
    except ValueError as exc:
        assert "encabezado" in str(exc)
    finally:
        empty.close()

    clash = tmp_path / "choque.xlsx"
    _workbook(clash, {"Datos": [["id"], [1]], "Datos!": [["id"], [2]]})
    collided = ExcelExtractor({"path": str(clash)})
    try:
        collided.connect()
        raise AssertionError("debía rechazar hojas con el mismo nombre")
    except ValueError as exc:
        assert "mismo nombre" in str(exc)

    filtered = tmp_path / "filtro.xlsx"
    _workbook(filtered, {"a": [["id"], [1]], "b": [["id"], [2]]})
    only_a = ExcelExtractor({"path": str(filtered), "sheets": ["a"]})
    only_a.connect()
    try:
        assert list(only_a.sheets) == ["a"]
    finally:
        only_a.close()
    unused = ExcelExtractor({"path": str(filtered), "sheets": ["no-existe"]})
    try:
        unused.connect()
        raise AssertionError("debía fallar si ninguna hoja coincide")
    except ValueError as exc:
        assert "utilizables" in str(exc)

    messy = tmp_path / "headers.xlsx"
    _workbook(
        messy,
        {"items": [["nombre", "nombre", ""], [True, "", "x"], ["Ana", "López", None]]},
    )
    fresh = ExcelExtractor({"path": str(messy)})
    assert list(fresh.extract("items"))[0]["nombre"] is True
    fresh.close()

    extractor = ExcelExtractor({"path": str(messy)})
    schema = extractor.discover_schema()
    assert [column.name for column in schema.tables[0].columns] == ["nombre", "nombre_2", "col_3"]
    assert schema.tables[0].columns[0].type == "boolean"
    rows = list(extractor.extract("items"))
    assert rows[0]["nombre"] is True
    assert rows[0]["nombre_2"] is None
    try:
        list(extractor.extract("no_existe"))
        raise AssertionError("debía fallar si la hoja no existe")
    except ValueError as exc:
        assert "no encontrada" in str(exc)
    finally:
        extractor.close()


def test_excel_checkpoint_skips_rows(tmp_path):
    from app.adapters.extractors.excel_extractor import ExcelExtractor

    path = tmp_path / "skip.xlsx"
    _workbook(path, {"items": [["id", "nombre"], [1, "Ana"], [2, "Beto"], [3, "Cia"]]})
    extractor = ExcelExtractor({"path": str(path)})
    extractor.connect()
    try:
        rows = list(extractor.extract("items", checkpoint={"skip": 2}))
    finally:
        extractor.close()
    assert rows == [{"id": 3, "nombre": "Cia"}]
